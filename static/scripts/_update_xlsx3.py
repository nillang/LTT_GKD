# -*- coding: utf-8 -*-
"""临时脚本：更新全文件注释完成后的进度表。"""
import openpyxl
from copy import copy

path = r'docs/LTT_GKD_项目进度跟踪_v3.xlsx'
wb = openpyxl.load_workbook(path)
D = '2026-09-11'

# ---------- 开发步骤 ----------
ws2 = wb['开发步骤']
steps = [
    ('十三·代码质量', '全文件中文注释补全（42 个 .kt 文件）', 'app/src/main/kotlin/**/*.kt', '已完成', 'KDoc+行内注释；数据层8+工具服务层13+UI/Theme7+已有14=42'),
]
base = ws2.max_row + 1
for i, (a, b, c, d, e) in enumerate(steps):
    rr = base + i
    for col, v in zip(range(1, 6), (a, b, c, d, e)):
        cell = ws2.cell(row=rr, column=col, value=v)
        cell._style = copy(ws2.cell(row=10, column=col)._style)

# ---------- 功能清单 ----------
ws = wb['功能清单']
# 更新 v5 改版行备注
ws.cell(row=63, column=8, value=(
    '首页电源圆脉冲/规则三Tab/卡片式设置/步骤编辑器/跳过记录页；'
    '单Activity；移除上传小FAB(导出在三点菜单)；'
    'Logger补clear()+LogViewer加返回/清空/自动刷新；'
    '全部42个.kt文件中文注释补全'
))

# ---------- 项目总览 ----------
ws4 = wb['项目总览']
ws4['B12'] = D

wb.save(path)
print('saved OK')
